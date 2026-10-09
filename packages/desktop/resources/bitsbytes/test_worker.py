import contextlib
import hashlib
import importlib.util
import io
import json
from pathlib import Path
import tempfile
import unittest
import wave
import os
import subprocess
import threading
import queue
from unittest.mock import patch
from types import SimpleNamespace

spec = importlib.util.spec_from_file_location("bbworker", Path(__file__).with_name("worker.py"))
worker = importlib.util.module_from_spec(spec)
spec.loader.exec_module(worker)
TEMPLATE = Path(r"\\10.10.1.3\AdminShare\documenten beheer administratief\2. Intakegesprekken\Intakegesprek blanco.docx")


class DurableWorkflow(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="bitsbytes-tests-")
        worker.ROOT = Path(self.temp.name)
        self.stdout = contextlib.redirect_stdout(io.StringIO())
        self.stdout.__enter__()
        self.addCleanup(self.stdout.__exit__, None, None, None)
        self.addCleanup(self.temp.cleanup)

    def test_intake_blank_fields_text_identifiers_and_idempotency(self):
        output = worker.ROOT / "output"
        output.mkdir()
        original = hashlib.sha256(TEMPLATE.read_bytes()).hexdigest()
        job = worker.create_job({"skill": "bits-en-bytes-intake", "template": str(TEMPLATE)})
        data = {"deelnemer_naam": "Voorbeeld Test", "deelnemer_telefoon": "0612345678", "deelnemer_bsn": "012345678", "achtergrond": "Eerste alinea.\n\nTweede alinea.", "afspraken": ["Twee dagen beoogd."]}
        args = {"id": job["id"], "data": data, "output": str(output)}
        first = worker.create_document(args)
        second = worker.create_document(args)
        self.assertEqual(first, second)
        self.assertEqual(len(list(output.glob("*.docx"))), 1)
        from docx import Document
        document = Document(first["path"])
        mapping = worker.helper().field_map(document)
        self.assertEqual(document.tables[0].rows[mapping["deelnemer_bsn"]].cells[1].text, "012345678")
        self.assertEqual(document.tables[0].rows[mapping["contactpersoon_naam"]].cells[1].text, "")
        self.assertEqual(hashlib.sha256(TEMPLATE.read_bytes()).hexdigest(), original)
        corrected = worker.create_document({**args, "data": {"deelnemer_telefoon": "0698765432"}, "revision": True})
        self.assertNotEqual(first["path"], corrected["path"])
        self.assertEqual(corrected["verified_field_count"], 13)
        revised = Document(corrected["path"])
        self.assertEqual(revised.tables[0].rows[mapping["deelnemer_naam"]].cells[1].text, "Voorbeeld Test")
        self.assertEqual(revised.tables[0].rows[mapping["achtergrond"]].cells[1].text, data["achtergrond"])
        self.assertEqual(revised.tables[0].rows[mapping["deelnemer_telefoon"]].cells[1].text, "0698765432")
        self.assertEqual(corrected, worker.create_document({**args, "data": {"deelnemer_telefoon": "0698765432"}, "revision": True}))

    def test_crash_after_document_write_is_recovered_and_parallel_calls_do_not_duplicate(self):
        output = worker.ROOT / "output"
        output.mkdir()
        job = worker.create_job({"template": str(TEMPLATE)})
        args = {"id": job["id"], "data": {"deelnemer_naam": "Crash Test"}, "output": str(output)}
        original = worker.update_job
        def crash(job_id, **values):
            if "artifact" in values:
                raise RuntimeError("simulated interruption after write")
            return original(job_id, **values)
        with patch.object(worker, "update_job", side_effect=crash):
            with self.assertRaisesRegex(RuntimeError, "simulated interruption"):
                worker.create_document(args)
        self.assertEqual(len(list(output.glob("*.docx"))), 1)
        self.assertTrue(worker.read_job(job["id"])["documentIntent"])
        recovered = worker.create_document(args)
        self.assertEqual(len(list(output.glob("*.docx"))), 1)
        with worker.concurrent.futures.ThreadPoolExecutor(max_workers=3) as pool:
            results = list(pool.map(worker.create_document, [args] * 3))
        self.assertTrue(all(row["path"] == recovered["path"] for row in results))

    def test_generic_template_partial_correction_preserves_prior_values(self):
        from docx import Document
        template = worker.ROOT / "generic.docx"
        document = Document()
        table = document.add_table(rows=2, cols=2)
        table.rows[0].cells[0].text = "Name"
        table.rows[1].cells[0].text = "Notes"
        document.save(template)
        job = worker.create_job({"template": str(template), "validateIntake": False})
        args = {"id": job["id"], "filename": "Test.docx", "output": str(worker.ROOT), "values": {"Name": "Example", "Notes": "Original"}}
        first = worker.create_generic(args)
        revised = worker.create_generic({**args, "values": {"Notes": "Corrected"}, "revision": True})
        self.assertNotEqual(first["path"], revised["path"])
        self.assertEqual(Document(revised["path"]).tables[0].rows[0].cells[1].text, "Example")
        self.assertEqual(revised["verified_field_count"], 2)

    def test_missing_output_fallback_and_retry(self):
        fallback = worker.ROOT / "local"
        result = worker.create_document({"template": str(TEMPLATE), "data": {}, "output": str(worker.ROOT / "missing"), "fallback": str(fallback)})
        self.assertTrue(result["used_fallback"])
        self.assertEqual(Path(result["path"]).name, "Intakegesprek deelnemer.docx")
        job = worker.create_job({})
        worker.update_job(job["id"], artifact=result)
        offline = worker.retry_save({"id": job["id"], "output": str(worker.ROOT / "offline")})
        self.assertEqual(result["path"], offline["path"])
        self.assertEqual(len(list(fallback.glob("*.docx"))), 1)
        destination = worker.ROOT / "server"
        destination.mkdir()
        saved = worker.retry_save({"id": job["id"], "output": str(destination)})
        self.assertFalse(saved["used_fallback"])
        self.assertEqual(Path(saved["path"]).read_bytes(), Path(result["path"]).read_bytes())
        self.assertEqual(saved, worker.retry_save({"id": job["id"], "output": str(destination)}))

    def test_template_change_and_unknown_fields_rejected(self):
        job = worker.create_job({"template": str(TEMPLATE)})
        Path(job["template"]).write_bytes(b"changed")
        with self.assertRaisesRegex(ValueError, "templateChanged"):
            worker.create_document({"id": job["id"], "data": {}, "output": str(worker.ROOT)})
        with self.assertRaises(ValueError):
            worker.helper().build_document(TEMPLATE, {"made_up_field": "invalid"})

    def test_long_chunk_manifest_and_interrupted_tail_recovery(self):
        job = worker.create_job({})
        payload = bytes(worker.RATE * 2 * worker.CHUNK_SECONDS)
        for _ in range(120):
            worker.add_chunk(job["id"], payload)
        worker.update_job(job["id"], recording=True, status="recording")
        (worker.job_dir(job["id"]) / "active.pcm").write_bytes(bytes(worker.RATE * 2 * 5))
        worker.recover()
        recovered = worker.read_job(job["id"])
        self.assertEqual(recovered["duration"], 3605)
        self.assertEqual(len(recovered["chunks"]), 121)
        self.assertFalse(recovered["recording"])
        worker.recover()
        self.assertEqual(worker.read_job(job["id"])["duration"], 3605)
        with self.assertRaises(ValueError):
            worker.job_dir("../../outside")

    def test_import_audio_and_share_probe_cleanup(self):
        source = worker.ROOT / "fixture.wav"
        with wave.open(str(source), "wb") as file:
            file.setnchannels(1); file.setsampwidth(2); file.setframerate(worker.RATE)
            file.writeframes(bytes(worker.RATE * 2 * 65))
        job = worker.import_audio({"path": str(source)})
        self.assertAlmostEqual(job["duration"], 65, places=3)
        self.assertEqual(len(job["chunks"]), 3)
        self.assertTrue(source.exists())
        result = worker.dispatch("share.test", {"path": str(worker.ROOT)})
        self.assertTrue(result["writable"])
        self.assertFalse(list(worker.ROOT.glob(".bitsbytes-check-*")))

    def test_empty_or_invalid_container_is_an_error_not_a_completed_recording(self):
        source = worker.ROOT / "empty.wav"
        with wave.open(str(source), "wb") as file:
            file.setnchannels(1); file.setsampwidth(2); file.setframerate(worker.RATE)
            file.writeframes(b"")
        with self.assertRaisesRegex(ValueError, "audioEmpty"):
            worker.import_audio({"path": str(source)})
        source = worker.ROOT / "invalid.m4a"
        source.write_bytes(b"invalid-container")
        with self.assertRaises(Exception):
            worker.import_audio({"path": str(source)})
        jobs = [json.loads(path.read_text()) for path in (worker.ROOT / "recordings").glob("*/job.json")]
        self.assertEqual(len(jobs), 2)
        self.assertTrue(all(job["status"] == "error" and job["importComplete"] is False for job in jobs))

    def test_crash_after_chunk_commit_does_not_duplicate_audio(self):
        job = worker.create_job({})
        payload = bytes(worker.RATE * 2 * 30)
        worker.add_chunk(job["id"], payload)
        (worker.job_dir(job["id"]) / "active-000000.pcm").write_bytes(payload)
        (worker.job_dir(job["id"]) / "active-000001.pcm").write_bytes(bytes(worker.RATE * 2 * 5))
        worker.update_job(job["id"], recording=True)
        worker.recover()
        self.assertEqual(worker.read_job(job["id"])["duration"], 35)

    def test_gpu_decode_failure_uses_cpu(self):
        gpu = SimpleNamespace(recognize=lambda _: (_ for _ in ()).throw(RuntimeError("CUDA failure")))
        cpu = SimpleNamespace(recognize=lambda _: "Dit is een test.")
        with patch.object(worker, "model", side_effect=lambda engine, target: cpu if target == "cpu" else gpu):
            result = worker.transcribe_chunk({"path":"fixture.wav","start":30,"duration":10}, "parakeet", {})
        self.assertEqual(result[0]["text"], "Dit is een test.")

    def test_interrupted_import_resumes_without_dropping_or_repeating_audio(self):
        source = worker.ROOT / "source.wav"
        with wave.open(str(source), "wb") as file:
            file.setnchannels(1); file.setsampwidth(2); file.setframerate(worker.RATE)
            file.writeframes(bytes(worker.RATE * 2 * 65))
        job = worker.create_job({})
        worker.add_chunk(job["id"], bytes(worker.RATE * 2 * 30))
        worker.update_job(job["id"], importSource=str(source), importComplete=False, status="interrupted")
        with self.assertRaisesRegex(ValueError, "importIncomplete"):
            worker.transcribe({"id":job["id"]})
        resumed = worker.dispatch("import.resume", {"id":job["id"]})
        self.assertAlmostEqual(resumed["duration"], 65)
        self.assertEqual(len(resumed["chunks"]), 3)
        self.assertTrue(resumed["importComplete"])

    def test_stdio_worker_import_does_not_deadlock_native_initialization(self):
        source = worker.ROOT / "stdio.wav"
        with wave.open(str(source), "wb") as file:
            file.setnchannels(1); file.setsampwidth(2); file.setframerate(worker.RATE)
            file.writeframes(bytes(worker.RATE * 2))
        child = subprocess.Popen([worker.sys.executable, "-u", str(Path(__file__).with_name("worker.py"))],
                                 stdin=subprocess.PIPE, stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True, encoding="utf-8",
                                 env={**os.environ, "BITSBYTES_DATA": str(worker.ROOT / "daemon"), "PYTHONUTF8": "1"})
        messages = queue.Queue()
        thread = threading.Thread(target=lambda: [messages.put(line) for line in child.stdout], daemon=True)
        thread.start()
        try:
            child.stdin.write(json.dumps({"id": "import", "method": "import", "args": {"path": str(source)}}) + "\n")
            child.stdin.flush()
            while True:
                message = json.loads(messages.get(timeout=15))
                if message.get("id") == "import":
                    self.assertNotIn("error", message)
                    self.assertEqual(message["result"]["duration"], 1)
                    break
        finally:
            child.stdin.close()
            try:
                child.wait(timeout=5)
            except subprocess.TimeoutExpired:
                child.kill(); child.wait()
            thread.join(timeout=5)
            child.stdout.close(); child.stderr.close()


if __name__ == "__main__":
    unittest.main(verbosity=2)
