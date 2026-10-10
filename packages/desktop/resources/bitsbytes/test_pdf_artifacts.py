"""Real PDF layout and image decoding regressions; no Office or app profile needed."""
import os
import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest

import fitz
from PIL import Image
from pypdf import PdfReader

import artifact_worker as worker


class PDFArtifactTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.folder = Path(tempfile.mkdtemp(prefix='bitsbytes-pdf-qa-'))
        os.environ.setdefault('MPLCONFIGDIR', str(cls.folder / 'matplotlib'))
        cls.image = cls.folder / 'source.png'
        Image.new('RGB', (1200, 300), '#255478').save(cls.image)

    def create(self, name, blocks, title=''):
        return Path(worker.create({'filename': name + '.pdf', 'title': title, 'blocks': blocks}, self.folder)['path'])

    def assert_margins(self, document):
        for page in document:
            for block in page.get_text('dict')['blocks']:
                box = fitz.Rect(block['bbox'])
                self.assertGreaterEqual(box.x0, 53)
                self.assertGreaterEqual(box.y0, 53)
                self.assertLessEqual(box.x1, page.rect.width - 53)
                self.assertLessEqual(box.y1, page.rect.height - 53)

    def test_all_blocks_unicode_content_and_rendering(self):
        literal = 'Één café: naïef, coöperatie, €25 — “afspraak”. Ελληνικά, Привет. <b>literal & safe</b>'
        path = self.create('complete', [
            {'type': 'heading', 'text': 'Afspraken', 'level': 1},
            {'type': 'paragraph', 'text': literal + '\nTweede regel.'},
            {'type': 'bullet', 'text': 'Een lijstpunt dat zichtbaar blijft.'},
            {'type': 'table', 'rows': [['Naam', 'Afspraak'], ['Zoë', 'Vrijdag om 10:00'], ['Renée', None]]},
            {'type': 'image', 'path': str(self.image), 'width': 25},
            {'type': 'pagebreak'},
            {'type': 'heading', 'text': 'Volgende pagina', 'level': 2},
            {'text': 'Einde van het document.'},
        ], title='Overzicht – 2026')
        reader = PdfReader(path)
        self.assertEqual(len(reader.pages), 2)
        first, second = [page.extract_text() for page in reader.pages]
        for value in ['Overzicht – 2026', 'Afspraken', literal, 'Tweede regel.', '•', 'Een lijstpunt', 'Zoë', 'Vrijdag om 10:00', 'Renée']:
            self.assertIn(value, first)
        self.assertNotIn('Volgende pagina', first)
        self.assertIn('Volgende pagina', second)
        with fitz.open(path) as document:
            self.assert_margins(document)
            spans = [span for block in document[0].get_text('dict')['blocks'] if block['type'] == 0
                     for line in block['lines'] for span in line['spans']]
            self.assertGreater(next(span['size'] for span in spans if span['text'] == 'Afspraken'),
                               next(span['size'] for span in spans if 'Één café' in span['text']))
            self.assertEqual(len(document[0].get_images()), 1)
            image = next(block for block in document[0].get_text('dict')['blocks'] if block['type'] == 1)
            box = fitz.Rect(image['bbox'])
            self.assertAlmostEqual(box.width / box.height, 4, places=4)
            for index, page in enumerate(document):
                pixmap = page.get_pixmap(matrix=fitz.Matrix(1.5, 1.5))
                pixmap.save(self.folder / f'complete-{index + 1}.png')
                self.assertTrue(any(value < 240 for value in pixmap.samples))
        self.assertEqual(worker.inspect(path)['pages'], 2)

    def test_wrapped_table_repeats_headers_and_keeps_every_row(self):
        rows = [['Identifier', 'Details']] + [[f'ROW-{index:03}', ('Useful readable details ' * 7) + f'END-{index:03}'] for index in range(85)]
        path = self.create('multipage-table', [{'type': 'table', 'rows': rows}])
        reader = PdfReader(path)
        self.assertGreater(len(reader.pages), 3)
        texts = [page.extract_text() for page in reader.pages]
        for text in texts:
            self.assertIn('Identifier', text)
            self.assertIn('Details', text)
        for index in range(85):
            self.assertEqual('\n'.join(texts).count(f'ROW-{index:03}'), 1)
            self.assertEqual('\n'.join(texts).count(f'END-{index:03}'), 1)
        with fitz.open(path) as document:
            self.assert_margins(document)
            document[1].get_pixmap().save(self.folder / 'table-continuation.png')

    def test_tall_row_and_unbroken_text_wrap_without_clipping(self):
        path = self.create('tall-row', [{'type': 'table', 'rows': [
            ['Reference', 'Long content'], ['item', ' '.join(f'word{index:04}' for index in range(550)) + ' FINALMARKER'],
        ]}, {'text': 'X' * 300 + ' ENDWORD'}])
        reader = PdfReader(path)
        self.assertGreater(len(reader.pages), 1)
        text = '\n'.join(page.extract_text() for page in reader.pages)
        self.assertIn('FINALMARKER', text)
        self.assertIn('ENDWORD', text)
        for index in range(550):
            self.assertEqual(text.count(f'word{index:04}'), 1)
        with fitz.open(path) as document:
            self.assert_margins(document)

    def test_tall_image_fits_page(self):
        source = self.folder / 'portrait.png'
        Image.new('RGB', (300, 1800), '#235478').save(source)
        path = self.create('portrait', [{'type': 'image', 'path': str(source), 'width': 20}])
        with fitz.open(path) as document:
            self.assertEqual(len(document), 1)
            self.assert_margins(document)
            box = fitz.Rect(next(block for block in document[0].get_text('dict')['blocks'] if block['type'] == 1)['bbox'])
            self.assertAlmostEqual(box.height / box.width, 6, places=4)

    def test_long_paragraph_flows_across_pages(self):
        words = ' '.join(f'paragraph{index:04}' for index in range(1800))
        path = self.create('long-paragraph', [{'type': 'heading', 'text': 'Complete content'}, {'text': words}])
        text = '\n'.join(page.extract_text() for page in PdfReader(path).pages)
        for index in range(1800):
            self.assertEqual(text.count(f'paragraph{index:04}'), 1)
        with fitz.open(path) as document:
            self.assertGreater(len(document), 3)
            self.assert_margins(document)

    def test_invalid_blocks_and_missing_glyphs_fail_before_publication(self):
        for block in [{'type': 'unsupported', 'text': 'Must not vanish'}, {'type': 'table', 'rows': []},
                      {'type': 'table', 'rows': ['not a row']}, {'type': 'heading', 'level': 'wrong'},
                      {'type': 'image', 'path': str(self.image), 'width': 0}, {'text': None}]:
            with self.subTest(block=block), self.assertRaises(ValueError):
                self.create('invalid', [block])
            self.assertFalse((self.folder / 'invalid.pdf').exists())
        with self.assertRaisesRegex(ValueError, 'bb.error.pdfUnsupportedText'):
            self.create('unsupported-glyph', [{'text': '\U0010ffff'}])
        self.assertFalse((self.folder / 'unsupported-glyph.pdf').exists())

    def test_word_blocks_still_have_native_structure(self):
        from docx import Document
        path = worker.create({'filename': 'word.docx', 'title': 'Word intact', 'blocks': [
            {'type': 'heading', 'text': 'Heading', 'level': 2}, {'type': 'bullet', 'text': 'Bullet'},
            {'type': 'table', 'rows': [['One', 'Two'], ['Three', 'Four']]},
            {'type': 'image', 'path': str(self.image), 'width': 2}, {'type': 'pagebreak'}, {'text': 'Last'},
        ]}, self.folder)['path']
        document = Document(path)
        self.assertEqual(document.tables[0].cell(1, 1).text, 'Four')
        self.assertEqual(next(p for p in document.paragraphs if p.text == 'Heading').style.name, 'Heading 2')
        self.assertEqual(next(p for p in document.paragraphs if p.text == 'Bullet').style.name, 'List Bullet')
        self.assertEqual(len(document.inline_shapes), 1)
        self.assertIn('w:type="page"', document.element.xml)
        with self.assertRaisesRegex(ValueError, 'bb.error.unsupportedArtifact'):
            worker.create({'filename': 'bad.docx', 'blocks': [{'type': 'not-supported'}]}, self.folder)

    @classmethod
    def tearDownClass(cls):
        print('PDF_QA_FOLDER=' + str(cls.folder))


class NormalizeImageTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory(prefix='bitsbytes-image-decode-')
        self.addCleanup(self.temporary.cleanup)
        self.folder = Path(self.temporary.name)

    def test_real_jpeg_webp_and_png_decode_to_png(self):
        for format in ['JPEG', 'WEBP', 'PNG']:
            with self.subTest(format=format):
                source = self.folder / ('source.' + format.lower())
                target = self.folder / ('normalized-' + format + '.png')
                Image.new('RGB', (80, 48), (30, 80, 120)).save(source, format=format)
                result = worker.normalize_image(source, target)
                self.assertEqual(result, {'width': 80, 'height': 48, 'sourceFormat': format})
                self.assertEqual(target.read_bytes()[:8], b'\x89PNG\r\n\x1a\n')
                with Image.open(target) as image:
                    self.assertEqual(image.mode, 'RGB')
                    self.assertEqual(image.size, (80, 48))
                    image.verify()
                with Image.open(target) as image:
                    self.assertTrue(all(abs(actual - expected) <= 3 for actual, expected in zip(image.getpixel((0, 0)), (30, 80, 120))))

    def test_worker_action_returns_decoded_metadata(self):
        source = self.folder / 'request.jpg'
        target = self.folder / 'result.png'
        Image.new('RGB', (40, 30), '#255478').save(source, format='JPEG')
        result = subprocess.run([sys.executable, '-B', str(Path(worker.__file__)), 'normalize_image'],
                                input=json.dumps({'path': str(source), 'output': str(target)}),
                                text=True, capture_output=True, check=True)
        self.assertEqual(json.loads(result.stdout), {'width': 40, 'height': 30, 'sourceFormat': 'JPEG'})
        self.assertEqual(target.read_bytes()[:8], b'\x89PNG\r\n\x1a\n')

    def test_transparency_is_preserved(self):
        for format in ['PNG', 'WEBP']:
            source = self.folder / ('transparent.' + format.lower())
            target = self.folder / ('transparent-' + format + '.png')
            Image.new('RGBA', (20, 16), (10, 20, 30, 80)).save(source, format=format)
            worker.normalize_image(source, target)
            with Image.open(target) as image:
                self.assertEqual(image.mode, 'RGBA')
                self.assertEqual(image.getpixel((0, 0))[3], 80)

    def test_corrupt_and_unsupported_inputs_do_not_create_output(self):
        source = self.folder / 'invalid'
        target = self.folder / 'output.png'
        for data in [b'<svg xmlns="http://www.w3.org/2000/svg"/>', b'not an image']:
            source.write_bytes(data)
            with self.assertRaisesRegex(ValueError, 'bb.error.imageOutput'):
                worker.normalize_image(source, target)
            self.assertFalse(target.exists())
        for format in ['GIF', 'BMP', 'TIFF']:
            Image.new('RGB', (20, 20)).save(source, format=format)
            with self.assertRaisesRegex(ValueError, 'bb.error.imageOutput'):
                worker.normalize_image(source, target)
            self.assertFalse(target.exists())
        for format in ['PNG', 'JPEG', 'WEBP']:
            Image.new('RGB', (80, 80)).save(source, format=format)
            source.write_bytes(source.read_bytes()[:-20])
            with self.assertRaisesRegex(ValueError, 'bb.error.imageOutput'):
                worker.normalize_image(source, target)
            self.assertFalse(target.exists())

    def test_overlarge_decoded_image_is_rejected(self):
        source = self.folder / 'large.png'
        Image.new('1', (10000, 4001)).save(source)
        target = self.folder / 'large-output.png'
        with self.assertRaisesRegex(ValueError, 'bb.error.imageOutput'):
            worker.normalize_image(source, target)
        self.assertFalse(target.exists())


if __name__ == '__main__':
    unittest.main()
