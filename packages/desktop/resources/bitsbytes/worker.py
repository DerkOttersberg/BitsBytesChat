"""Local durable recording, speech and Word worker. JSON lines over stdio only."""
from __future__ import annotations
import argparse
import concurrent.futures
import hashlib
import importlib.util
import json
import os
from pathlib import Path
import queue
import shutil
import sys
import threading
import time
import uuid
import wave
import math
import array
import subprocess
import base64
import re
from fractions import Fraction
from io import BytesIO

RATE = 16000
CHUNK_SECONDS = 30
LOCK = threading.RLock()
PRINT_LOCK = threading.Lock()
POOL = concurrent.futures.ThreadPoolExecutor(max_workers=4)
ASR_LOCK = threading.Lock()
DOCUMENT_LOCKS = {}
MODELS = {}
CANCEL = {}
RECORD = None
ROOT = Path(os.environ.get("BITSBYTES_DATA", Path.home() / "AppData/Local/BitsBytesChat"))
ASSETS = Path(__file__).parent
OPENROUTER_STT_URL = "https://openrouter.ai/api/v1/audio/transcriptions"
AUDIO_DEFAULT_MODEL = "elevenlabs/scribe-v2"
DIARIZATION_MODELS = {"elevenlabs/scribe-v2", "elevenlabs/scribe-v2-medical"}
DIARIZATION_UPLOAD_LIMIT = 25_000_000


def emit(value):
    with PRINT_LOCK:
        print(json.dumps(value, ensure_ascii=False), flush=True)


def atomic_json(path, value):
    path.parent.mkdir(parents=True, exist_ok=True)
    temp = path.with_suffix(f".{uuid.uuid4().hex}.tmp")
    temp.write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding="utf-8")
    os.replace(temp, path)


def job_dir(job_id):
    if not isinstance(job_id, str) or not len(job_id) == 32 or any(c not in "0123456789abcdef" for c in job_id):
        raise ValueError("bb.error.invalidJob")
    return ROOT / "recordings" / job_id


def read_job(job_id):
    with LOCK:
        return json.loads((job_dir(job_id) / "job.json").read_text(encoding="utf-8"))


def save_job(job):
    with LOCK:
        job["updated"] = time.time()
        atomic_json(job_dir(job["id"]) / "job.json", job)
    emit({"event": "job", "job": job})


def update_job(job_id, **values):
    with LOCK:
        job = read_job(job_id)
        job.update(values)
        save_job(job)
        return job


def helper():
    spec = importlib.util.spec_from_file_location("intake_docx", ASSETS / "skills/bits-en-bytes-intake/scripts/intake_docx.py")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def create_job(args):
    job = {"id": uuid.uuid4().hex, "title": args.get("title") or time.strftime("Recording %Y-%m-%d %H.%M"),
           "created": time.time(), "updated": time.time(), "status": "ready", "chunks": [],
           "duration": 0, "skill": args.get("skill", ""), "engine": args.get("engine", "parakeet"),
           "segments": [], "transcript": "", "error": "", "recording": False}
    if job["engine"] == "openrouter":
        job.update(audioModel=args.get("audioModel", AUDIO_DEFAULT_MODEL), audioLanguage=args.get("audioLanguage", ""),
                   audioDiarization=args.get("audioDiarization", False), speakerLabel=args.get("speakerLabel", "speaker_{{number}}"))
        if job["audioDiarization"] and job["audioModel"] not in DIARIZATION_MODELS:
            raise ValueError("bb.error.audioDiarizationUnsupported")
    directory = job_dir(job["id"])
    directory.mkdir(parents=True)
    if args.get("template"):
        snapshot = directory / "template.docx"
        try:
            shutil.copyfile(args["template"], snapshot)
        except OSError as exc:
            raise ValueError("bb.error.templateMissing") from exc
        if args.get("validateIntake", True):
            helper().build_document(snapshot, {})
        else:
            from docx import Document
            Document(snapshot)
        job["template"] = str(snapshot)
        job["templateHash"] = hashlib.sha256(snapshot.read_bytes()).hexdigest()
    save_job(job)
    return job


def add_chunk(job_id, payload):
    with LOCK:
        job = read_job(job_id)
        index = len(job["chunks"])
        path = job_dir(job_id) / f"audio-{index:06d}.wav"
        with wave.open(str(path), "wb") as file:
            file.setnchannels(1)
            file.setsampwidth(2)
            file.setframerate(RATE)
            file.writeframes(payload)
        duration = len(payload) / (RATE * 2)
        job["chunks"].append({"path": str(path), "start": job["duration"], "duration": duration, "done": False})
        job["duration"] += duration
        save_job(job)


def capture_format(sounddevice, device):
    try:
        info = sounddevice.query_devices(device, "input")
    except sounddevice.PortAudioError as error:
        raise ValueError("bb.error.audioMicrophone") from error
    rates = list(dict.fromkeys([RATE, int(info["default_samplerate"])]))
    channels = [1] + ([2] if info["max_input_channels"] >= 2 else [])
    for rate in rates:
        for channel in channels:
            try:
                sounddevice.check_input_settings(device=device, channels=channel, dtype="int16", samplerate=rate)
                return rate, channel
            except sounddevice.PortAudioError:
                continue
    raise ValueError("bb.error.audioMicrophone")


class CapturePCM:
    def __init__(self, rate, channels):
        import av
        self.av = av
        self.rate = rate
        self.channels = channels
        self.samples = 0
        self.resampler = av.AudioResampler(format="s16", layout="mono", rate=RATE)

    def convert(self, data):
        if data is None:
            frames = self.resampler.resample(None)
        else:
            count = len(data) // (2 * self.channels)
            frame = self.av.AudioFrame(format="s16", layout="mono" if self.channels == 1 else "stereo", samples=count)
            frame.planes[0].update(data)
            frame.sample_rate = self.rate
            frame.time_base = Fraction(1, self.rate)
            frame.pts = self.samples
            self.samples += count
            frames = self.resampler.resample(frame)
        return b"".join(bytes(frame.planes[0])[:frame.samples * 2] for frame in frames)


class Recorder:
    def __init__(self, job_id, device):
        import sounddevice
        rate, channels = capture_format(sounddevice, device)
        self.pcm = CapturePCM(rate, channels)
        self.job_id = job_id
        self.paused = threading.Event()
        self.ending = threading.Event()
        self.data = queue.Queue(maxsize=256)
        self.error = ""
        self.seconds = 0
        self.thread = threading.Thread(target=self.write, daemon=True)
        self.thread.start()
        try:
            self.stream = sounddevice.RawInputStream(samplerate=rate, channels=channels, dtype="int16", device=device,
                                                     blocksize=max(1, rate // 10), callback=self.callback)
            self.stream.start()
        except Exception as error:
            if hasattr(self, "stream"):
                self.stream.close()
            self.ending.set()
            self.thread.join(timeout=15)
            if isinstance(error, sounddevice.PortAudioError):
                raise ValueError("bb.error.audioMicrophone") from error
            raise

    def callback(self, data, frames, info, status):
        if self.paused.is_set() or self.ending.is_set():
            return
        if status:
            self.error = str(status)
        try:
            self.data.put_nowait(bytes(data))
        except queue.Full:
            self.error = "bb.error.recordingOverflow"

    def write(self):
        payload = bytearray()
        last_meter = 0
        index = len(read_job(self.job_id)["chunks"])
        partial = job_dir(self.job_id) / f"active-{index:06d}.pcm"
        journal = partial.open("wb")
        try:
            while not self.ending.is_set() or not self.data.empty():
                try:
                    data = self.data.get(timeout=0.2)
                except queue.Empty:
                    continue
                data = self.pcm.convert(data)
                journal.write(data)
                journal.flush()
                if time.monotonic() - last_meter > 0.2:
                    samples = array.array("h", data)
                    level = math.sqrt(sum(value * value for value in samples) / max(1, len(samples))) / 32768
                    emit({"event": "meter", "level": level})
                    last_meter = time.monotonic()
                payload.extend(data)
                self.seconds += len(data) / (RATE * 2)
                if time.monotonic() - getattr(self, "progress_at", 0) >= 1:
                    self.progress_at = time.monotonic()
                    emit({"event": "duration", "id": self.job_id, "duration": self.seconds})
                if len(payload) >= CHUNK_SECONDS * RATE * 2:
                    add_chunk(self.job_id, bytes(payload))
                    payload.clear()
                    journal.close()
                    partial.unlink()
                    index += 1
                    partial = job_dir(self.job_id) / f"active-{index:06d}.pcm"
                    journal = partial.open("wb")
            tail = self.pcm.convert(None)
            if tail:
                journal.write(tail)
                journal.flush()
                payload.extend(tail)
            if payload:
                add_chunk(self.job_id, bytes(payload))
            journal.close()
            partial.unlink(missing_ok=True)
        except Exception as exc:
            self.error = str(exc)
            self.ending.set()
            update_job(self.job_id, status="error", error=self.error)
        finally:
            journal.close()

    def stop(self):
        self.stream.stop()
        self.stream.close()
        self.ending.set()
        self.thread.join(timeout=15)
        if self.thread.is_alive():
            raise RuntimeError("bb.error.recordingFlush")
        job = read_job(self.job_id)
        status = job["status"] if job["status"] in ("error", "cancelled", "transcribing") else "recorded"
        return update_job(self.job_id, recording=False, status=status, error=self.error or job.get("error", ""))


def import_audio(args):
    import av
    job = read_job(args["id"]) if args.get("id") else create_job(args)
    source = Path(job.get("importSource") or args["path"])
    fingerprint = {"size": source.stat().st_size, "modified": source.stat().st_mtime_ns}
    if job.get("importFingerprint") and job["importFingerprint"] != fingerprint:
        raise ValueError("bb.error.audioChanged")
    update_job(job["id"], importSource=str(source), importFingerprint=fingerprint, importComplete=False)
    update_job(job["id"], status="importing")
    buffer = bytearray()
    skip = round(job["duration"] * RATE) * 2
    container = None
    resampler = av.AudioResampler(format="s16", layout="mono", rate=RATE)
    def append(audio):
        nonlocal skip
        payload = audio.to_ndarray().tobytes()
        count = min(skip, len(payload))
        skip -= count
        buffer.extend(payload[count:])
    try:
        container = av.open(str(source))
        for frame in container.decode(audio=0):
            for audio in resampler.resample(frame):
                append(audio)
                while len(buffer) >= CHUNK_SECONDS * RATE * 2:
                    size = CHUNK_SECONDS * RATE * 2
                    add_chunk(job["id"], bytes(buffer[:size]))
                    del buffer[:size]
        for audio in resampler.resample(None):
            append(audio)
        if buffer:
            add_chunk(job["id"], bytes(buffer))
        if not read_job(job["id"])["duration"]:
            raise ValueError("bb.error.audioEmpty")
    except Exception as exc:
        update_job(job["id"], status="error", error=str(exc))
        raise
    finally:
        if container:
            container.close()
    return update_job(job["id"], status="recorded", importComplete=True)


def model(engine, acceleration="auto"):
    key = (engine, acceleration)
    if key in MODELS:
        return MODELS[key]
    if engine == "parakeet":
        import onnx_asr
        import onnxruntime
        providers = onnxruntime.get_available_providers()
        target = ["CUDAExecutionProvider", "CPUExecutionProvider"] if acceleration != "cpu" and "CUDAExecutionProvider" in providers else ["CPUExecutionProvider"]
        try:
            value = onnx_asr.load_model("nemo-parakeet-tdt-0.6b-v3", quantization="int8", providers=target)
        except Exception:
            value = onnx_asr.load_model("nemo-parakeet-tdt-0.6b-v3", quantization="int8", providers=["CPUExecutionProvider"])
        # VAD prevents silence hallucinations and bounds model windows on long recordings.
        value = value.with_vad(onnx_asr.load_vad("silero"))
        MODELS[key] = value
    if engine == "whisper":
        from faster_whisper import WhisperModel
        device = "cpu"
        if acceleration != "cpu":
            try:
                import ctranslate2
                if ctranslate2.get_cuda_device_count():
                    device = "cuda"
            except Exception:
                pass
        try:
            MODELS[key] = WhisperModel("large-v3-turbo", device=device, compute_type="int8_float16" if device == "cuda" else "int8",
                                      download_root=str(ROOT / "models/whisper"))
        except Exception:
            MODELS[key] = WhisperModel("large-v3-turbo", device="cpu", compute_type="int8", download_root=str(ROOT / "models/whisper"))
    if key not in MODELS:
        raise ValueError("bb.error.invalidEngine")
    return MODELS[key]


def diarization_audio(recording):
    """One speech-quality MP3 request retains speaker identity across capture chunks."""
    import av
    import numpy
    encoded = BytesIO()
    with av.open(encoded, mode="w", format="mp3") as container:
        stream = container.add_stream("libmp3lame", rate=RATE)
        stream.layout = "mono"
        stream.bit_rate = 32000
        for chunk in recording["chunks"]:
            with wave.open(chunk["path"], "rb") as audio:
                if (audio.getnchannels(), audio.getsampwidth(), audio.getframerate()) != (1, 2, RATE):
                    raise ValueError("bb.error.audioResponse")
                while payload := audio.readframes(RATE * 10):
                    frame = av.AudioFrame.from_ndarray(numpy.frombuffer(payload, dtype="<i2").reshape(1, -1),
                                                       format="s16", layout="mono")
                    frame.sample_rate = RATE
                    for packet in stream.encode(frame):
                        container.mux(packet)
                    if encoded.tell() > DIARIZATION_UPLOAD_LIMIT:
                        raise ValueError("bb.error.audioDiarizationLimit")
        for packet in stream.encode(None):
            container.mux(packet)
    if len(encoded.getbuffer()) > DIARIZATION_UPLOAD_LIMIT:
        raise ValueError("bb.error.audioDiarizationLimit")
    return encoded.getvalue()


def speaker_segments(result, recording):
    rows = result.get("words") or result.get("segments")
    if not isinstance(rows, list) or len(rows) > 100_000:
        if not result["text"].strip():
            return []
        raise ValueError("bb.error.audioDiarizationResponse")
    turns = []
    labelled = False
    for row in rows:
        if not isinstance(row, dict):
            raise ValueError("bb.error.audioDiarizationResponse")
        text = row.get("word", row.get("text"))
        if not isinstance(text, str):
            raise ValueError("bb.error.audioDiarizationResponse")
        if not text.strip() or row.get("type") in ("spacing", "audio_event"):
            continue
        start, end = row.get("start"), row.get("end")
        if (any(isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value) for value in (start, end))
                or not 0 <= start <= end <= recording["duration"] + 2):
            raise ValueError("bb.error.audioDiarizationResponse")
        speaker = row.get("speaker", row.get("speaker_label", row.get("speaker_id")))
        if type(speaker) is int and 0 <= speaker < 32:
            speaker = f"speaker_{speaker}"
        else:
            if isinstance(speaker, str) and re.fullmatch(r"(?:[0-9]|[12][0-9]|3[01])", speaker):
                speaker = "speaker_" + speaker
            if speaker is not None and (not isinstance(speaker, str) or not re.fullmatch(r"speaker_(?:[0-9]|[12][0-9]|3[01])", speaker)):
                raise ValueError("bb.error.audioDiarizationResponse")
        labelled = labelled or speaker is not None
        if turns and turns[-1].get("speaker") == speaker:
            if start < turns[-1]["start"]:
                raise ValueError("bb.error.audioDiarizationResponse")
            turns[-1]["words"].append(text.strip())
            turns[-1]["end"] = max(turns[-1]["end"], end)
            continue
        turn = {"start": recording["start"] + start, "end": recording["start"] + end, "words": [text.strip()]}
        if speaker is not None:
            turn["speaker"] = speaker
        turns.append(turn)
    if result["text"].strip() and not labelled:
        raise ValueError("bb.error.audioDiarizationResponse")
    for turn in turns:
        turn["text"] = re.sub(r"\s+([,.!?;:])", r"\1", " ".join(turn.pop("words")))
    return turns


def save_transcript(job):
    names = job.get("speakerNames", {})
    lines = []
    for segment in job["segments"]:
        if not segment["text"].strip():
            continue
        speaker = segment.get("speaker")
        name = names.get(speaker) if speaker else None
        if speaker and not name:
            name = job.get("speakerLabel", "speaker_{{number}}").replace("{{number}}", str(int(speaker.split("_")[1]) + 1))
        lines.append((name + ": " if name else "") + segment["text"].strip())
    job["transcript"] = "\n".join(lines)
    path = job_dir(job["id"]) / "transcript.txt"
    temporary = path.with_suffix(f".{uuid.uuid4().hex}.tmp")
    try:
        temporary.write_text(job["transcript"], encoding="utf-8")
        os.replace(temporary, path)
    finally:
        temporary.unlink(missing_ok=True)


def openrouter_request(chunk, args, endpoint=OPENROUTER_STT_URL):
    import httpx
    diarization = args.get("audioDiarization", False)
    if diarization and args["audioModel"] not in DIARIZATION_MODELS:
        raise ValueError("bb.error.audioDiarizationUnsupported")
    payload = {"model": args["audioModel"], "input_audio": {
        "data": base64.b64encode(diarization_audio(chunk) if diarization else Path(chunk["path"]).read_bytes()).decode("ascii"),
        "format": "mp3" if diarization else "wav"},
        "response_format": "verbose_json" if diarization else "json"}
    if diarization:
        payload.update(timestamp_granularities=["word"], provider={"options": {"elevenlabs": {"diarize": True}}})
    if args.get("audioLanguage"):
        payload["language"] = args["audioLanguage"]
    try:
        # Stream a bounded response; never include upstream bodies, credentials or audio in errors.
        with httpx.stream("POST", endpoint, headers={"Authorization": "Bearer " + args["key"],
                          "X-Title": "Bits & Bytes Chat"}, json=payload, timeout=180 if diarization else 60) as response:
            errors = {401: "audioKeyRejected", 403: "audioKeyRejected", 402: "audioCredits",
                      429: "audioQuota", 400: "audioModelRejected", 404: "audioModelRejected", 422: "audioModelRejected",
                      408: "audioDiarizationUncertain" if diarization else "audioUncertain",
                      504: "audioDiarizationUncertain" if diarization else "audioUncertain"}
            if response.status_code != 200:
                raise RuntimeError("bb.error." + errors.get(response.status_code, "audioProviderError"))
            body = bytearray()
            for part in response.iter_bytes():
                body.extend(part)
                if len(body) > (32 * 1024 * 1024 if diarization else 1024 * 1024):
                    raise RuntimeError("bb.error.audioResponse")
    except httpx.ConnectError:
        raise RuntimeError("bb.error.audioNetwork") from None
    except httpx.HTTPError:
        # A timed-out POST might have been processed. Leave this chunk uncommitted for explicit retry.
        raise RuntimeError("bb.error.audioDiarizationUncertain" if diarization else "bb.error.audioUncertain") from None
    try:
        result = json.loads(body)
    except (ValueError, UnicodeError):
        raise RuntimeError("bb.error.audioResponse") from None
    if not isinstance(result, dict) or not isinstance(result.get("text"), str) or len(result["text"]) > (2 * 1024 * 1024 if diarization else 262144):
        raise RuntimeError("bb.error.audioResponse")
    if diarization:
        return speaker_segments(result, chunk)
    return [{"start": chunk["start"], "end": chunk["start"] + chunk["duration"], "text": result["text"].strip()}]


class AudioCancelled(Exception):
    pass


def transcribe_cloud(chunk, args):
    if not isinstance(args.get("key"), str) or not args["key"].strip():
        raise ValueError("bb.error.audioKey")
    cancellation = args.get("_cancellation") or threading.Event()
    if cancellation.is_set():
        raise AudioCancelled()
    # The key travels over a private pipe, never process arguments, disk or job events.
    child = subprocess.Popen([sys.executable, str(ASSETS / "worker.py"), "--openrouter-transcribe"],
                             stdin=subprocess.PIPE, stdout=subprocess.PIPE, stderr=subprocess.PIPE,
                             env={**os.environ, "PYTHONUTF8": "1"},
                             creationflags=subprocess.CREATE_NO_WINDOW if os.name == "nt" else 0)
    request = json.dumps({"chunk": chunk, "args": {key: args.get(key) for key in ("key", "audioModel", "audioLanguage", "audioDiarization")},
                          "endpoint": OPENROUTER_STT_URL}).encode("utf-8")
    deadline = time.monotonic() + (215 if args.get("audioDiarization") else 65)
    try:
        while True:
            if cancellation.is_set():
                raise AudioCancelled()
            if time.monotonic() >= deadline:
                raise RuntimeError("bb.error.audioDiarizationUncertain" if args.get("audioDiarization") else "bb.error.audioUncertain")
            try:
                output, _ = child.communicate(input=request, timeout=0.1)
                break
            except subprocess.TimeoutExpired:
                request = None
        if cancellation.is_set():
            raise AudioCancelled()
        try:
            result = json.loads(output)
        except (ValueError, UnicodeError):
            raise RuntimeError("bb.error.audioResponse") from None
        if child.returncode or not isinstance(result, dict):
            raise RuntimeError("bb.error.audioResponse")
        if result.get("error"):
            raise RuntimeError(result["error"])
        if not isinstance(result.get("segments"), list):
            raise RuntimeError("bb.error.audioResponse")
        return result["segments"]
    finally:
        if child.poll() is None:
            child.kill()
        child.communicate()


def transcribe_chunk(chunk, engine, args):
    if engine == "openrouter":
        return transcribe_cloud(chunk, {"audioModel": AUDIO_DEFAULT_MODEL, "audioLanguage": "", **args})
    if engine == "groq":
        raise ValueError("bb.error.audioRetired")
    acceleration = args.get("acceleration", "auto")
    def run_local(target):
        value = model(engine, target)
        if engine == "whisper":
            segments, _ = value.transcribe(chunk["path"], language="nl", beam_size=5, vad_filter=True, condition_on_previous_text=False)
            return [{"start": chunk["start"] + row.start, "end": chunk["start"] + row.end, "text": row.text.strip()} for row in segments]
        return value.recognize(chunk["path"])
    try:
        result = run_local(acceleration)
    except ModuleNotFoundError:
        raise ValueError("bb.error.audioLocalSetup") from None
    except Exception:
        if acceleration == "cpu":
            raise
        # CUDA/cuDNN errors can surface on the first decode, after model construction.
        MODELS.pop((engine, acceleration), None)
        result = run_local("cpu")
        MODELS[(engine, acceleration)] = model(engine, "cpu")
    if engine == "whisper":
        return result
    if isinstance(result, str):
        return [{"start": chunk["start"], "end": chunk["start"] + chunk["duration"], "text": result}]
    return [{"start": chunk["start"] + float(getattr(row, "start", 0)),
             "end": chunk["start"] + float(getattr(row, "end", chunk["duration"])), "text": row.text} for row in result]


def transcribe(args):
    job_id = args["id"]
    if read_job(job_id).get("importComplete") is False:
        raise ValueError("bb.error.importIncomplete")
    cancellation = CANCEL.setdefault(job_id, threading.Event())
    cancellation.clear()
    started = time.time()
    # Capture remains chunked and recoverable. One finalized provider request gives
    # the entire recording a single speaker namespace; per-chunk IDs cannot be joined safely.
    while read_job(job_id).get("audioDiarization") and read_job(job_id).get("recording") and not cancellation.is_set():
        time.sleep(0.1)
    if cancellation.is_set():
        return update_job(job_id, status="cancelled")
    with ASR_LOCK:
        if cancellation.is_set():
            return update_job(job_id, status="cancelled")
        update_job(job_id, status="transcribing", error="")
        try:
            job = read_job(job_id)
            if not job.get("recording") and not job["chunks"]:
                raise ValueError("bb.error.audioEmpty")
            if job.get("audioDiarization") and not job.get("diarizationComplete"):
                if not job["chunks"]:
                    raise ValueError("bb.error.audioEmpty")
                segments = transcribe_cloud({"start": 0, "duration": job["duration"], "chunks": job["chunks"]},
                    {**args, "_cancellation": cancellation, "audioDiarization": True,
                     "audioModel": job.get("audioModel", AUDIO_DEFAULT_MODEL), "audioLanguage": job.get("audioLanguage", "")})
                if cancellation.is_set():
                    return update_job(job_id, status="cancelled")
                with LOCK:
                    job = read_job(job_id)
                    job.update(segments=segments, diarizationComplete=True)
                    for chunk in job["chunks"]:
                        chunk.update(done=True, segments=[])
                    save_transcript(job)
                    save_job(job)
            if job.get("audioDiarization"):
                return update_job(job_id, status="transcribed", processingSeconds=time.time() - started)
            while not cancellation.is_set():
                job = read_job(job_id)
                chunk = next((row for row in job["chunks"] if not row.get("done")), None)
                if chunk is None:
                    if job.get("recording"):
                        time.sleep(0.5)
                        continue
                    break
                segments = transcribe_chunk(chunk, job["engine"], {**args, "_cancellation": cancellation,
                    "audioModel": job.get("audioModel", AUDIO_DEFAULT_MODEL), "audioLanguage": job.get("audioLanguage", "")})
                if cancellation.is_set():
                    break
                with LOCK:
                    job = read_job(job_id)
                    index = next(i for i, row in enumerate(job["chunks"]) if row["path"] == chunk["path"])
                    job["chunks"][index]["done"] = True
                    job["chunks"][index]["segments"] = segments
                    job["segments"] = [item for row in job["chunks"] for item in row.get("segments", [])]
                    save_transcript(job)
                    save_job(job)
        except AudioCancelled:
            cancellation.set()
        except Exception as exc:
            update_job(job_id, status="error", error=str(exc))
            raise
        return update_job(job_id, status="cancelled" if cancellation.is_set() else "transcribed", processingSeconds=time.time() - started)


def export_preview(args):
    import pythoncom
    import win32com.client
    pythoncom.CoInitialize()
    word = None
    document = None
    destination = ROOT / "previews" / (hashlib.sha256(Path(args["path"]).read_bytes()).hexdigest() + ".pdf")
    destination.parent.mkdir(parents=True, exist_ok=True)
    # Resolve the created directory before passing it across COM. Windows Store
    # launchers can virtualize AppData; Word must receive the physical path.
    destination = destination.parent.resolve() / destination.name
    try:
        if destination.is_file():
            return {"path": str(destination)}
        word = win32com.client.DispatchEx("Word.Application")
        word.Visible = False
        word.DisplayAlerts = 0
        word.AutomationSecurity = 3
        document = word.Documents.Open(str(Path(args["path"]).resolve()), ReadOnly=True, AddToRecentFiles=False)
        document.ExportAsFixedFormat(str(destination), 17)
        if not destination.is_file() or not destination.read_bytes().startswith(b"%PDF-"):
            raise IOError("bb.error.verifySave")
        return {"path": str(destination)}
    finally:
        if document is not None:
            document.Close(False)
        if word is not None:
            word.Quit()
        pythoncom.CoUninitialize()


def preview(args):
    # Isolate COM from the recording worker. A Word add-in cannot block Stop or all jobs.
    try:
        result = subprocess.run([sys.executable, str(ASSETS / "worker.py"), "--preview", str(Path(args["path"]).resolve())],
                                capture_output=True, text=True, encoding="utf-8", timeout=90,
                                creationflags=subprocess.CREATE_NO_WINDOW if os.name == "nt" else 0)
    except subprocess.TimeoutExpired as exc:
        raise RuntimeError("bb.error.wordPreview") from exc
    if result.returncode:
        raise RuntimeError("bb.error.wordPreview")
    return json.loads(result.stdout)


def document_lock(job_id):
    with LOCK:
        return DOCUMENT_LOCKS.setdefault(job_id or "standalone", threading.RLock())


def document_template(job, args):
    template = Path(job["template"] if job and job.get("template") else args["template"])
    try:
        digest = hashlib.sha256(template.read_bytes()).hexdigest()
    except OSError as exc:
        raise ValueError("bb.error.templateMissing") from exc
    if job and job.get("templateHash") and digest != job["templateHash"]:
        raise ValueError("bb.error.templateChanged")
    return template, digest


def finish_document(job_id):
    """Recover a completed external write before trying another filename.

    The exact verified bytes and every candidate path are journalled locally
    before writing to a share. A crash after the write cannot create a second
    document on replay. Exclusive creation never overwrites another intake.
    """
    job = read_job(job_id)
    intent = job["documentIntent"]
    payload = (job_dir(job_id) / intent["payload"]).read_bytes()
    if hashlib.sha256(payload).hexdigest() != intent["hash"]:
        raise ValueError("bb.error.verifySave")

    def complete(target):
        artifact = {**intent["result"], "saved": True, "path": target["path"], "used_fallback": target["fallback"]}
        update_job(job_id, artifact=artifact, documentData=intent["data"], documentSignature=intent["signature"], documentIntent=None, status="complete")
        return artifact

    for target in intent["targets"]:
        try:
            if Path(target["path"]).read_bytes() == payload:
                return complete(target)
        except OSError:
            pass

    for folder, fallback in [(Path(intent["output"]), False), (Path(intent["fallback"]), True)]:
        if fallback:
            existing = intent.get("existingFallback")
            if existing and Path(existing).read_bytes() == payload:
                return complete({"path": existing, "fallback": True})
            folder.mkdir(parents=True, exist_ok=True)
        for number in range(1, 1001):
            filename = intent["filename"] if number == 1 else f"{Path(intent['filename']).stem} ({number}).docx"
            destination = folder / filename
            try:
                if destination.exists():
                    continue
                target = {"path": str(destination), "fallback": fallback}
                intent["targets"].append(target)
                update_job(job_id, documentIntent=intent)
                helper().write_new(payload, destination)
            except FileExistsError:
                continue
            except OSError:
                if fallback:
                    raise
                break
            return complete(target)
    raise ValueError("bb.error.verifySave")


def save_document(job, payload, filename, data, signature, args, result=None, stamp=True):
    if not job:
        used_fallback = False
        try:
            path = helper().write_version(payload, Path(args["output"]), filename)
        except OSError:
            folder = Path(args.get("fallback") or ROOT / "documents")
            folder.mkdir(parents=True, exist_ok=True)
            path = helper().write_version(payload, folder, filename)
            used_fallback = True
        return {**(result or {}), "saved": True, "path": str(path), "used_fallback": used_fallback}
    operation = uuid.uuid4().hex
    if stamp:
        # Distinguish otherwise identical intakes from different recording jobs.
        from docx import Document
        document = Document(BytesIO(payload))
        document.core_properties.identifier = f"bitsbytes:{job['id']}:{operation}"
        memory = BytesIO()
        document.save(memory)
        payload = memory.getvalue()
    local = job_dir(job["id"]) / f"document-{operation}.docx"
    helper().write_new(payload, local)
    intent = {"payload": local.name, "hash": hashlib.sha256(payload).hexdigest(), "signature": signature,
              "filename": filename, "data": data, "result": result or {}, "targets": [],
              "output": args["output"], "fallback": args.get("fallback") or str(ROOT / "documents"), "existingFallback": args.get("existingFallback")}
    update_job(job["id"], documentIntent=intent)
    return finish_document(job["id"])


def create_document(args):
    job_id = args.get("id")
    with document_lock(job_id):
        job = read_job(job_id) if job_id else None
        if job and job.get("documentIntent"):
            finish_document(job_id)
            job = read_job(job_id)
        if job and job.get("artifact") and not args.get("revision"):
            return job["artifact"]
        template, digest = document_template(job, args)
        previous = job.get("documentData", {}) if job and args.get("revision") else {}
        if job and args.get("revision") and job.get("artifact") and not previous:
            # Migrate jobs created by an earlier app version without losing fields.
            from docx import Document
            document = Document(job["artifact"]["path"])
            previous = {key: document.tables[0].rows[row].cells[1].text for key, row in helper().field_map(document).items()}
        payload, prepared, count = helper().build_document(template, {**previous, **args["data"]})
        signature = hashlib.sha256(json.dumps([digest, prepared], sort_keys=True, ensure_ascii=False).encode()).hexdigest()
        if job and job.get("artifact") and job.get("documentSignature") == signature:
            return job["artifact"]
        return save_document(job, payload, helper().document_name(prepared), prepared, signature, args,
                             {"verified_field_count": count, "blank_field_count": sum(not value for value in prepared.values())})


def retry_save(args):
    with document_lock(args["id"]):
        job = read_job(args["id"])
        if job.get("documentIntent"):
            return finish_document(job["id"])
        artifact = job["artifact"]
        if not artifact.get("used_fallback") and Path(artifact["path"]).parent == Path(args["output"]):
            return artifact
        payload = Path(artifact["path"]).read_bytes()
        return save_document(job, payload, Path(artifact["path"]).name, job.get("documentData", {}), job.get("documentSignature", ""),
                             {**args, "fallback": str(Path(artifact["path"]).parent), "existingFallback": artifact["path"]}, artifact, stamp=False)


def create_generic(args):
    from docx import Document
    job_id = args.get("id")
    with document_lock(job_id):
        job = read_job(job_id) if job_id else None
        if job and job.get("documentIntent"):
            finish_document(job_id)
            job = read_job(job_id)
        if job and job.get("artifact") and not args.get("revision"):
            return job["artifact"]
        template, digest = document_template(job, args)
        document = Document(template)
        values = {**(job.get("documentData", {}) if job and args.get("revision") else {}), **args["values"]}
        remaining = dict(values)
        expected = []
        for t, table in enumerate(document.tables):
            for i, row in enumerate(table.rows):
                if len(row.cells) < 2 or row.cells[0]._tc is row.cells[1]._tc:
                    continue
                key = str(i) if str(i) in remaining else row.cells[0].text.strip()
                if key in remaining:
                    value = str(remaining.pop(key))
                    helper().replace_answer(row.cells[1], value)
                    expected.append((t, i, value))
        if remaining:
            raise ValueError("Unmapped template fields: " + ", ".join(remaining))
        memory = BytesIO()
        document.save(memory)
        payload = memory.getvalue()
        reopened = Document(BytesIO(payload))
        if any(reopened.tables[t].rows[i].cells[1].text != value for t, i, value in expected):
            raise ValueError("bb.error.verifySave")
        filename = Path(args["filename"]).name
        if filename != args["filename"] or not filename.lower().endswith(".docx") or any(c in filename for c in '<>:"/\\|?*'):
            raise ValueError("bb.error.invalidRequest")
        signature = hashlib.sha256(json.dumps([digest, filename, values], sort_keys=True, ensure_ascii=False).encode()).hexdigest()
        if job and job.get("artifact") and job.get("documentSignature") == signature:
            return job["artifact"]
        return save_document(job, payload, filename, values, signature, args, {"verified_field_count": len(expected)})


def dispatch(method, args):
    global RECORD
    if method == "health":
        # The interpreter can predate a first-use install. Activate the newly
        # created component directory before checking its imports.
        import site
        site.addsitedir(str(ROOT / "runtime" / "worker" / "Lib" / "site-packages"))
        importlib.invalidate_caches()
        try:
            for name in ("sounddevice", "av", "httpx"):
                importlib.import_module(name)
            ready = True
        except (ImportError, OSError):
            ready = False
        return {"python": sys.version.split()[0], "root": str(ROOT), "engines": ["parakeet", "whisper", "openrouter"], "audioReady": ready}
    if method == "devices":
        import sounddevice
        return [{"id": i, "name": row["name"]} for i, row in enumerate(sounddevice.query_devices()) if row["max_input_channels"] > 0]
    if method == "jobs":
        return sorted([json.loads(path.read_text(encoding="utf-8")) for path in (ROOT / "recordings").glob("*/job.json")], key=lambda row: row["created"], reverse=True)
    if method == "job":
        return read_job(args["id"])
    if method == "template.refresh":
        with document_lock(args["id"]):
            job = read_job(args["id"])
            if job.get("recording") or job.get("documentIntent"):
                raise ValueError("bb.error.activeJob")
            staging = job_dir(job["id"]) / "template-new.docx"
            try:
                shutil.copyfile(args["template"], staging)
            except OSError as exc:
                raise ValueError("bb.error.templateMissing") from exc
            if args.get("validateIntake", True):
                helper().build_document(staging, {})
            else:
                from docx import Document
                Document(staging)
            snapshot = job_dir(job["id"]) / "template.docx"
            os.replace(staging, snapshot)
            return update_job(job["id"], template=str(snapshot), templateHash=hashlib.sha256(snapshot.read_bytes()).hexdigest(), error="")
    if method == "job.update":
        allowed = {"sessionID", "engine", "skill", "audioModel", "audioLanguage", "audioDiarization", "speakerLabel"}
        values = {key: value for key, value in args["values"].items() if key in allowed}
        with LOCK:
            job = read_job(args["id"])
            changed = (values.get("engine", job["engine"]) != job["engine"] or
                       values.get("engine", job["engine"]) == "openrouter" and any(
                           key in values and values[key] != job.get(key, default)
                           for key, default in (("audioModel", AUDIO_DEFAULT_MODEL), ("audioLanguage", ""), ("audioDiarization", False))))
            if "audioDiarization" in values and type(values["audioDiarization"]) is not bool:
                raise ValueError("bb.error.invalidSettings")
            if values.get("audioDiarization", job.get("audioDiarization")) and values.get("audioModel", job.get("audioModel")) not in DIARIZATION_MODELS:
                raise ValueError("bb.error.audioDiarizationUnsupported")
            if changed:
                if job.get("recording") or job.get("status") == "transcribing":
                    raise ValueError("bb.error.activeJob")
                for chunk in job["chunks"]:
                    chunk["done"] = False
                    chunk.pop("segments", None)
                job.update(values, segments=[], transcript="", status="recorded", error="")
                job.pop("diarizationComplete", None)
                job.pop("speakerNames", None)
                (job_dir(job["id"]) / "transcript.txt").write_text("", encoding="utf-8")
                save_job(job)
                return job
            return update_job(args["id"], **values)
    if method == "job.speakers":
        with LOCK:
            job = read_job(args["id"])
            if job.get("recording") or job["status"] == "transcribing":
                raise ValueError("bb.error.activeJob")
            speakers = {segment.get("speaker") for segment in job["segments"] if segment.get("speaker")}
            names = args.get("names")
            if (not isinstance(names, dict) or any(speaker not in speakers or not isinstance(name, str)
                    or len(name.strip()) > 80 or any(ord(char) < 32 or ord(char) == 127 for char in name)
                    for speaker, name in names.items())):
                raise ValueError("bb.error.audioSpeakerName")
            job["speakerNames"] = {speaker: name.strip() for speaker, name in names.items() if name.strip()}
            save_transcript(job)
            save_job(job)
            return job
    if method == "import":
        return import_audio(args)
    if method == "import.resume":
        job = read_job(args["id"])
        return import_audio(args) if job.get("importComplete") is False else job
    if method == "record.start":
        if RECORD:
            raise ValueError("bb.error.alreadyRecording")
        job = create_job(args)
        update_job(job["id"], recording=True, status="recording")
        try:
            RECORD = Recorder(job["id"], args.get("device"))
        except Exception as exc:
            update_job(job["id"], recording=False, status="error", error=str(exc))
            raise
        return read_job(job["id"])
    if method in ("record.pause", "record.resume", "record.stop"):
        if not RECORD:
            raise ValueError("bb.error.noRecording")
        if method == "record.pause":
            RECORD.paused.set()
            return update_job(RECORD.job_id, status="paused")
        if method == "record.resume":
            RECORD.paused.clear()
            return update_job(RECORD.job_id, status="recording")
        result = RECORD.stop()
        RECORD = None
        return result
    if method == "transcribe":
        return transcribe(args)
    if method == "transcribe.cancel":
        CANCEL.setdefault(args["id"], threading.Event()).set()
        return {"cancelled": True}
    if method == "model.download":
        try:
            with ASR_LOCK:
                model(args["engine"], "cpu")
        except ModuleNotFoundError:
            raise ValueError("bb.error.audioLocalSetup") from None
        return {"ready": True}
    if method == "docx.create":
        return create_document(args)
    if method == "docx.preview":
        return preview(args)
    if method == "docx.retry":
        return retry_save(args)
    if method == "docx.inspect":
        return helper().inspect_template(Path(args["path"]))
    if method == "docx.labels":
        from docx import Document
        return [[{"row": i, "label": row.cells[0].text, "columns": len(row.cells)} for i, row in enumerate(table.rows)] for table in Document(args["path"]).tables]
    if method == "docx.generic":
        return create_generic(args)
    if method == "share.test":
        folder = Path(args["path"])
        probe = folder / (".bitsbytes-check-" + uuid.uuid4().hex + ".tmp")
        try:
            with probe.open("xb") as file:
                file.write(b"BitsBytesChat storage check")
            if probe.read_bytes() != b"BitsBytesChat storage check":
                raise IOError("bb.error.verifySave")
            return {"readable": True, "writable": True}
        finally:
            probe.unlink(missing_ok=True)
    if method == "job.delete":
        with document_lock(args["id"]):
            job = read_job(args["id"])
            if job.get("recording") or job.get("status") == "transcribing":
                raise ValueError("bb.error.activeJob")
            shutil.rmtree(job_dir(args["id"]))
        return {"deleted": True}
    raise ValueError("bb.error.unknownMethod")


def handle(request):
    try:
        result = dispatch(request["method"], request.get("args") or {})
        emit({"id": request["id"], "result": result})
    except Exception as exc:
        emit({"id": request["id"], "error": str(exc)})


def recover():
    for path in (ROOT / "recordings").glob("*/job.json"):
        job = json.loads(path.read_text(encoding="utf-8"))
        if job.get("recording") or job["status"] in ("transcribing", "importing"):
            for journal in sorted(path.parent.glob("active-*.pcm")):
                index = int(journal.stem.removeprefix("active-"))
                if index == len(read_job(job["id"])["chunks"]) and journal.stat().st_size:
                    payload = journal.read_bytes()
                    add_chunk(job["id"], payload[:len(payload) // 2 * 2])
                # A journal whose index is already committed cannot duplicate audio.
                journal.unlink()
            partial = path.parent / "active.pcm"
            if partial.is_file() and partial.stat().st_size:
                add_chunk(job["id"], partial.read_bytes())
                partial.unlink()
            update_job(job["id"], recording=False, status="interrupted", error="bb.error.interrupted")


if __name__ == "__main__":
    if len(sys.argv) == 2 and sys.argv[1] == "--openrouter-transcribe":
        try:
            request = json.load(sys.stdin)
            print(json.dumps({"segments": openrouter_request(request["chunk"], request["args"], request["endpoint"])}, ensure_ascii=False))
        except Exception as exc:
            message = str(exc)
            print(json.dumps({"error": message if message.startswith("bb.error.audio") else "bb.error.audioResponse"}))
        sys.exit(0)
    if len(sys.argv) == 3 and sys.argv[1] == "--preview":
        print(json.dumps(export_preview({"path": sys.argv[2]})))
        sys.exit(0)
    # Initialize NumPy's Windows native runtime on the main thread. Loading it
    # lazily from an AV/ASR job while the main thread blocks on stdin can deadlock.
    import numpy
    ROOT.mkdir(parents=True, exist_ok=True)
    os.environ["HF_HOME"] = str(ROOT / "models/huggingface")
    recover()
    for line in sys.stdin:
        try:
            request = json.loads(line)
            POOL.submit(handle, request)
        except Exception:
            emit({"error": "bb.error.invalidRequest"})
    if RECORD:
        RECORD.stop()
    for cancellation in CANCEL.values():
        cancellation.set()
    POOL.shutdown(wait=True)
