#!/usr/bin/env python3
"""Inspect or fill the known Bits en Bytes intake table without changing labels.

Requires python-docx. Missing or unusable identifiers remain blank. Does not
transcribe, overwrite files, or accept a changed template without inspection.
Sensitive answer values are never printed.
"""

from __future__ import annotations

import argparse
from copy import deepcopy
from io import BytesIO
import json
from pathlib import Path
import re
import sys

from docx import Document


DEFAULT_TEMPLATE = Path(r"\\10.10.1.3\AdminShare\documenten beheer administratief\2. Intakegesprekken\Intakegesprek blanco.docx")
DEFAULT_DIRECTORY = Path(r"\\10.10.1.3\AdminShare\documenten beheer administratief\2. Intakegesprekken")
EXPECTED_LABELS = (
    "Datum", "", "Intakegesprek met", "Naam", "Adres", "Telefoonnummer",
    "Geboortedatum", "Geslacht", "BSN", "", "Contactpersoon", "Naam",
    "Telefoonnummer", "",
    "Vorm van financiering - ZIN (WMO/WLZ) - PGB (WLZ) - Onderaannemerschap",
    "Achtergrond", "Affiniteit met ICT", "Gemaakte afspraken",
)
GLOBAL_FIELDS = {
    "datum": "datum",
    "vorm van financiering - zin (wmo/wlz) - pgb (wlz) - onderaannemerschap": "financiering",
    "achtergrond": "achtergrond",
    "affiniteit met ict": "ict_affiniteit",
    "gemaakte afspraken": "afspraken",
}
PERSON_FIELDS = {
    "naam": "naam", "adres": "adres", "telefoonnummer": "telefoon",
    "geboortedatum": "geboortedatum", "geslacht": "geslacht", "bsn": "bsn",
}
UNKNOWN_VALUES = frozenset({
    "Niet besproken", "Onbekend", "Nog te bevestigen", "Niet aangeleverd", "N.v.t.",
})


class IntakeError(ValueError):
    pass


def normalized(text: str) -> str:
    return " ".join(text.split()).casefold()


def field_map(document):
    if len(document.tables) != 1:
        raise IntakeError("Afwijkend template: verwacht één intake-tabel.")
    rows = document.tables[0].rows
    if any(len(row.cells) != 2 for row in rows):
        raise IntakeError("Afwijkend template: verwacht twee kolommen per rij.")
    if any(row.cells[0]._tc is row.cells[1]._tc for row in rows):
        raise IntakeError("Afwijkend template: samengevoegde antwoordcellen.")
    if any(cell.tables for row in rows for cell in row.cells):
        raise IntakeError("Afwijkend template: geneste tabellen.")
    if document.element.xpath(".//w:sdt | .//w:ffData | .//w:del | .//w:ins"):
        raise IntakeError("Afwijkend template: formuliervelden of bijgehouden wijzigingen.")
    labels = [normalized(row.cells[0].text) for row in rows]
    if labels != [normalized(label) for label in EXPECTED_LABELS]:
        raise IntakeError("Afwijkende veldindeling: inspecteer het actuele template met documenttools.")
    fields = {}
    section = None
    for index, label in enumerate(labels):
        if label == "intakegesprek met":
            section = "deelnemer"
        elif label == "contactpersoon":
            section = "contactpersoon"
        elif label in GLOBAL_FIELDS:
            fields[GLOBAL_FIELDS[label]] = index
        elif label in PERSON_FIELDS and section:
            key = f"{section}_{PERSON_FIELDS[label]}"
            if key in fields:
                raise IntakeError("Dubbel veld in dezelfde persoonssectie.")
            fields[key] = index
    if len(fields) != 13:
        raise IntakeError("Onvolledige veldkoppeling.")
    return fields


def inspect_template(template: Path):
    document = Document(template)
    result = {
        "table_count": len(document.tables),
        "labels": [[row.cells[0].text for row in table.rows] for table in document.tables],
    }
    try:
        result["field_rows_zero_based"] = field_map(document)
        result["supported"] = True
    except IntakeError as exc:
        result["supported"] = False
        result["reason"] = str(exc)
    return result


def prepare_data(data, fields):
    if not isinstance(data, dict):
        raise IntakeError("Gegevens moeten een JSON-object zijn; onbekende velden mogen ontbreken.")
    if set(data) - set(fields):
        raise IntakeError("Onbekende veldsleutel: controleer de sleutels in word-invullen.md.")
    prepared = {}
    for key in fields:
        value = data.get(key)
        if value is None:
            value = ""
        if key == "afspraken" and isinstance(value, list):
            if any(not isinstance(item, str) for item in value):
                raise IntakeError("afspraken moet een lijst van teksten zijn.")
            value = [item.strip() for item in value if item.strip()]
        elif key in {"deelnemer_bsn", "deelnemer_telefoon", "contactpersoon_telefoon"} and not isinstance(value, str):
            # Numeric input could already have lost leading zeroes. Do not guess.
            value = ""
        elif not isinstance(value, str):
            raise IntakeError(f"{key} moet tekst zijn, of leeg blijven.")
        else:
            value = value.strip()
        prepared[key] = value
    bsn = prepared["deelnemer_bsn"]
    if bsn and bsn not in UNKNOWN_VALUES:
        joined = "".join(bsn.split())
        prepared["deelnemer_bsn"] = joined if re.fullmatch(r"[0-9]{9}", joined) else ""
    return prepared


def replace_answer(cell, value):
    texts = ["- " + item for item in value] if isinstance(value, list) else [value]
    first = cell.paragraphs[0]
    paragraph_properties = deepcopy(first._p.pPr)
    run_properties = deepcopy(first.runs[0]._r.rPr) if first.runs else None
    for paragraph in list(cell.paragraphs)[1:]:
        paragraph._p.getparent().remove(paragraph._p)
    first.clear()
    for index, text in enumerate(texts):
        if index == 0:
            paragraph = first
        else:
            paragraph = cell.add_paragraph(style=first.style)
            if paragraph_properties is not None:
                if paragraph._p.pPr is not None:
                    paragraph._p.remove(paragraph._p.pPr)
                paragraph._p.insert(0, deepcopy(paragraph_properties))
        run = paragraph.add_run(text)
        if run_properties is not None:
            run._r.insert(0, deepcopy(run_properties))


def build_document(template: Path, data):
    document = Document(template)
    fields = field_map(document)
    prepared = prepare_data(data, fields)
    table = document.tables[0]
    expected = {}
    for key, row_index in fields.items():
        value = prepared[key]
        replace_answer(table.rows[row_index].cells[1], value)
        expected[key] = "\n".join("- " + item for item in value) if isinstance(value, list) else value
    buffer = BytesIO()
    document.save(buffer)
    buffer.seek(0)
    reopened = Document(buffer)
    reopened_fields = field_map(reopened)
    for key, row_index in reopened_fields.items():
        if reopened.tables[0].rows[row_index].cells[1].text != expected[key]:
            raise IntakeError(f"Controle van het ingevulde veld {key} is mislukt.")
    return buffer.getvalue(), prepared, len(fields)


def write_new(payload: bytes, output: Path):
    if not output.parent.is_dir():
        raise FileNotFoundError("De uitvoermap bestaat niet.")
    with output.open("xb") as destination:
        destination.write(payload)
    if output.read_bytes() != payload:
        raise OSError("De opgeslagen inhoud wijkt af van de gecontroleerde inhoud.")


def fill_template(template: Path, data, output: Path):
    if template.resolve() == output.resolve():
        raise IntakeError("Het blanco template mag niet worden overschreven.")
    if output.suffix.casefold() != ".docx":
        raise IntakeError("Het doelbestand moet .docx zijn.")
    if output.exists():
        raise IntakeError("Doelbestand bestaat al; gebruik create voor een afzonderlijke versie.")
    payload, prepared, count = build_document(template, data)
    write_new(payload, output)
    return {"saved": True, "verified_field_count": count, "blank_field_count": sum(not value for value in prepared.values())}


def document_name(prepared):
    name = prepared["deelnemer_naam"] or "deelnemer"
    name = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "", name).strip().rstrip(". ")
    return f"Intakegesprek {name or 'deelnemer'}.docx"


def write_version(payload: bytes, directory: Path, filename: str):
    stem = Path(filename).stem
    for version in range(1, 1001):
        name = filename if version == 1 else f"{stem} ({version}).docx"
        output = directory / name
        try:
            write_new(payload, output)
        except FileExistsError:
            continue
        return output
    raise IntakeError("Te veel gelijknamige versies; kies een andere bestemming.")


def create_intake(template: Path, data, output_dir: Path, fallback_dir: Path | None = None):
    payload, prepared, count = build_document(template, data)
    filename = document_name(prepared)
    used_fallback = False
    try:
        output = write_version(payload, output_dir, filename)
    except OSError:
        if fallback_dir is None:
            raise
        fallback_dir.mkdir(parents=True, exist_ok=True)
        output = write_version(payload, fallback_dir, filename)
        used_fallback = True
    return {
        "saved": True,
        "path": str(output),
        "used_fallback": used_fallback,
        "verified_field_count": count,
        "blank_field_count": sum(not value for value in prepared.values()),
    }


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    commands = parser.add_subparsers(dest="command", required=True)
    inspect = commands.add_parser("inspect", help="Toon structuur en labels, zonder antwoorden.")
    inspect.add_argument("--template", required=True, type=Path)
    fill = commands.add_parser("fill", help="Vul een nieuw lokaal .docx-bestand in.")
    fill.add_argument("--template", required=True, type=Path)
    fill.add_argument("--data", required=True, type=Path)
    fill.add_argument("--output", required=True, type=Path)
    create = commands.add_parser("create", help="Maak en bewaar direct; lege velden, versies en lokale terugval zijn toegestaan.")
    create.add_argument("--template", type=Path, default=DEFAULT_TEMPLATE)
    create.add_argument("--data", required=True, type=Path)
    create.add_argument("--output-dir", type=Path, default=DEFAULT_DIRECTORY)
    create.add_argument("--fallback-dir", type=Path)
    args = parser.parse_args()
    try:
        if args.command == "inspect":
            result = inspect_template(args.template)
        else:
            data = json.loads(args.data.read_text(encoding="utf-8-sig"))
            if args.command == "create":
                result = create_intake(args.template, data, args.output_dir, args.fallback_dir)
            else:
                result = fill_template(args.template, data, args.output)
    except IntakeError as exc:
        print(str(exc), file=sys.stderr)
        return 2
    except json.JSONDecodeError:
        print("Het gegevensbestand is geen geldig JSON-object.", file=sys.stderr)
        return 2
    except Exception as exc:
        print(f"Bestandsbewerking niet afgerond ({type(exc).__name__}); controleer bestanden en toegang.", file=sys.stderr)
        return 2
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    sys.exit(main())
