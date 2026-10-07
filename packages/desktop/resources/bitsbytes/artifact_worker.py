"""Independent artifact subprocess; never shares the recording worker or Office instances."""
import csv
import hashlib
import json
import os
import shutil
from pathlib import Path
import sys
import zipfile

MIME = {'.docx': 'application/vnd.openxmlformats-officedocument.wordprocessingml.document', '.xlsx': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', '.pptx': 'application/vnd.openxmlformats-officedocument.presentationml.presentation', '.pdf': 'application/pdf', '.csv': 'text/csv', '.html': 'text/html', '.htm': 'text/html', '.png': 'image/png', '.jpg': 'image/jpeg', '.jpeg': 'image/jpeg', '.svg': 'image/svg+xml'}
TEXT = {'.txt', '.md', '.json', '.js', '.ts', '.tsx', '.jsx', '.css', '.py', '.ps1', '.sql', '.xml', '.yaml', '.yml', '.toml', '.ini', '.sh', '.c', '.h', '.cpp', '.cs', '.java', '.rs', '.go', '.html', '.htm', '.csv', '.svg'}


def validate(path):
    path = Path(path).resolve(strict=True)
    if not path.is_file() or not path.stat().st_size or path.stat().st_size > 256 * 1024 * 1024:
        raise ValueError('bb.error.invalidDocument')
    ext = path.suffix.lower()
    if ext in {'.docx', '.xlsx', '.pptx'}:
        with zipfile.ZipFile(path) as archive:
            part = {'.docx': 'word/document.xml', '.xlsx': 'xl/workbook.xml', '.pptx': 'ppt/presentation.xml'}[ext]
            if part not in archive.namelist() or archive.testzip():
                raise ValueError('bb.error.invalidDocument')
    elif ext == '.pdf':
        from pypdf import PdfReader
        if path.read_bytes()[:5] != b'%PDF-' or not len(PdfReader(path).pages):
            raise ValueError('bb.error.invalidDocument')
    elif ext in TEXT:
        path.read_text(encoding='utf-8-sig')
    elif ext in {'.png', '.jpg', '.jpeg'}:
        from PIL import Image
        with Image.open(path) as image:
            image.verify()
    else:
        raise ValueError('bb.error.unsupportedArtifact')
    kind = {'.docx': 'document', '.pptx': 'presentation', '.xlsx': 'spreadsheet', '.pdf': 'pdf', '.csv': 'spreadsheet', '.html': 'html', '.htm': 'html', '.png': 'image', '.jpg': 'image', '.jpeg': 'image', '.svg': 'image'}.get(ext, 'code')
    return {'path': str(path), 'kind': kind, 'mime': MIME.get(ext, 'text/plain'), 'size': path.stat().st_size, 'hash': hashlib.sha256(path.read_bytes()).hexdigest()}


def inspect(path):
    info = validate(path)
    path = Path(info['path'])
    if path.suffix.lower() == '.docx':
        from docx import Document
        document = Document(path)
        info.update(paragraphs=[p.text for p in document.paragraphs], tables=[[[c.text for c in r.cells] for r in t.rows] for t in document.tables])
    elif path.suffix.lower() == '.pptx':
        from pptx import Presentation
        deck = Presentation(path)
        info.update(slides=[{'text': [s.text for s in slide.shapes if s.has_text_frame]} for slide in deck.slides])
    elif path.suffix.lower() == '.xlsx':
        from openpyxl import load_workbook
        book = load_workbook(path, read_only=True, data_only=False)
        info.update(sheets=[{'name': s.title, 'rows': s.max_row, 'columns': s.max_column, 'sample': [[str(v) if v is not None else '' for v in row] for row in s.iter_rows(max_row=min(s.max_row, 30), max_col=min(s.max_column, 30), values_only=True)]} for s in book])
        book.close()
    elif path.suffix.lower() == '.pdf':
        from pypdf import PdfReader
        pages = PdfReader(path).pages
        info.update(pages=len(pages), text=[p.extract_text() for p in pages])
    else:
        info['text'] = path.read_text(encoding='utf-8-sig') if path.suffix.lower() in TEXT else ''
    return info


def office(path, output):
    import pythoncom
    from win32com.client import DispatchEx
    pythoncom.CoInitialize()
    application = None
    document = None
    ext = path.suffix.lower()
    try:
        if ext == '.docx':
            application = DispatchEx('Word.Application')
            application.Visible = False
            application.DisplayAlerts = 0
            application.AutomationSecurity = 3
            document = application.Documents.Open(str(path), ReadOnly=True, AddToRecentFiles=False)
            document.ExportAsFixedFormat(str(output), 17)
        elif ext == '.pptx':
            application = DispatchEx('PowerPoint.Application')
            application.AutomationSecurity = 3
            document = application.Presentations.Open(str(path), ReadOnly=True, Untitled=False, WithWindow=False)
            document.SaveAs(str(output), 32)
        elif ext == '.xlsx':
            application = DispatchEx('Excel.Application')
            application.Visible = False
            application.DisplayAlerts = False
            application.AutomationSecurity = 3
            application.AskToUpdateLinks = False
            document = application.Workbooks.Open(str(path), UpdateLinks=0, ReadOnly=False)
            application.CalculateFullRebuild()
            document.Save()
            errors = []
            for sheet in document.Worksheets:
                values = sheet.UsedRange.Value
                if isinstance(values, tuple):
                    for r, row in enumerate(values):
                        if isinstance(row, tuple):
                            errors.extend({'sheet': sheet.Name, 'row': r + 1, 'column': c + 1, 'error': value} for c, value in enumerate(row) if isinstance(value, int) and value in {-2146826281, -2146826265, -2146826259, -2146826246, -2146826252, -2146826273, -2146826288})
            document.ExportAsFixedFormat(0, str(output))
            return {'pdf': str(output), 'formulaErrors': errors}
    finally:
        if document is not None:
            if ext == '.docx':
                document.Close(False)
            elif ext == '.xlsx':
                document.Close(False)
            else:
                document.Close()
        if application is not None:
            application.Quit()
        pythoncom.CoUninitialize()
    return {'pdf': str(output)}


def render(path, folder):
    path = Path(path).resolve(strict=True)
    folder = Path(folder)
    folder.mkdir(parents=True, exist_ok=True)
    output = folder / (hashlib.sha256(path.read_bytes()).hexdigest() + '.pdf')
    if path.suffix.lower() == '.pdf':
        return {'pdf': str(path)}
    if path.suffix.lower() in {'.docx', '.pptx', '.xlsx'}:
        if path.suffix.lower() == '.xlsx':
            cached = output.with_suffix('.xlsx')
            shutil.copyfile(path, cached)
            return office(cached, output)
        return office(path, output)
    return {'pdf': ''}


def grid(path, sheet=None, offset=0, limit=200):
    path = Path(path)
    if path.suffix.lower() == '.csv':
        with path.open(encoding='utf-8-sig', newline='') as stream:
            rows = list(csv.reader(stream))
        return {'sheets': ['CSV'], 'sheet': 'CSV', 'rows': rows[offset:offset + limit], 'total': len(rows), 'offset': offset}
    from openpyxl import load_workbook
    book = load_workbook(path, read_only=True, data_only=True)
    formulas = load_workbook(path, read_only=True, data_only=False)
    selected = sheet if sheet in book.sheetnames else book.sheetnames[0]
    source = book[selected]
    values = [[str(value) if value is not None else '' for value in row] for row in source.iter_rows(min_row=offset + 1, max_row=min(source.max_row, offset + limit), max_col=min(source.max_column, 100), values_only=True)]
    formula_values = [[str(value) if value is not None else '' for value in row] for row in formulas[selected].iter_rows(min_row=offset + 1, max_row=min(source.max_row, offset + limit), max_col=min(source.max_column, 100), values_only=True)]
    result = {'sheets': book.sheetnames, 'sheet': selected, 'rows': values, 'formulas': formula_values, 'total': source.max_row, 'columns': source.max_column, 'offset': offset}
    book.close()
    formulas.close()
    return result


def create(spec, folder):
    folder = Path(folder)
    folder.mkdir(parents=True, exist_ok=True)
    filename = spec['filename']
    if Path(filename).name != filename or not filename:
        raise ValueError('bb.error.invalidFilename')
    path = folder / filename
    ext = path.suffix.lower()
    if ext == '.docx':
        from docx import Document
        from docx.shared import Inches, Pt
        document = Document()
        section = document.sections[0]
        section.top_margin = section.bottom_margin = Inches(.75)
        document.styles['Normal'].font.name = 'Calibri'
        document.styles['Normal'].font.size = Pt(11)
        if spec.get('title'):
            document.add_heading(spec['title'], 0)
        for block in spec.get('blocks', []):
            kind = block.get('type', 'paragraph')
            if kind == 'heading':
                document.add_heading(block.get('text', ''), min(4, max(1, int(block.get('level', 1)))))
            elif kind == 'table':
                rows = block.get('rows', [])
                if not rows:
                    continue
                table = document.add_table(rows=0, cols=max(map(len, rows)))
                table.style = 'Light Shading Accent 1'
                for row in rows:
                    cells = table.add_row().cells
                    for index, value in enumerate(row):
                        cells[index].text = str(value)
            elif kind == 'image':
                document.add_picture(block['path'], width=Inches(min(6, float(block.get('width', 5)))))
            elif kind == 'pagebreak':
                document.add_page_break()
            else:
                document.add_paragraph(block.get('text', ''), style='List Bullet' if kind == 'bullet' else None)
        document.save(path)
    elif ext == '.xlsx':
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        from openpyxl.utils import get_column_letter
        book = Workbook()
        book.remove(book.active)
        for definition in spec.get('sheets', [{'name': 'Sheet1', 'rows': []}]):
            sheet = book.create_sheet(definition.get('name', 'Sheet1')[:31])
            for row in definition.get('rows', []):
                sheet.append(row)
            sheet.freeze_panes = 'A2'
            if sheet.max_row > 1:
                sheet.auto_filter.ref = sheet.dimensions
            for cell in sheet[1]:
                cell.font = Font(bold=True, color='FFFFFF')
                cell.fill = PatternFill('solid', fgColor='255478')
            for column in sheet.columns:
                letter = get_column_letter(column[0].column)
                sheet.column_dimensions[letter].width = min(55, max(12, max(len(str(cell.value or '')) for cell in column) + 2))
        book.save(path)
    elif ext == '.pptx':
        from pptx import Presentation
        from pptx.util import Inches, Pt
        deck = Presentation()
        deck.slide_width = Inches(13.333)
        deck.slide_height = Inches(7.5)
        for definition in spec.get('slides', []):
            slide = deck.slides.add_slide(deck.slide_layouts[1])
            slide.shapes.title.text = definition.get('title', '')
            slide.shapes.title.text_frame.paragraphs[0].font.size = Pt(32)
            frame = slide.placeholders[1].text_frame
            frame.clear()
            for index, text in enumerate(definition.get('bullets', [])):
                paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
                paragraph.text = str(text)
                paragraph.font.size = Pt(22)
            if definition.get('notes'):
                slide.notes_slide.notes_text_frame.text = definition['notes']
        deck.save(path)
    elif ext == '.pdf':
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
        from reportlab.lib.styles import getSampleStyleSheet
        from xml.sax.saxutils import escape
        styles = getSampleStyleSheet()
        blocks = [Paragraph(escape(spec.get('title', '')), styles['Title']), Spacer(1, 12)]
        blocks.extend(Paragraph(escape(block.get('text', '')).replace('\n', '<br/>'), styles['BodyText']) for block in spec.get('blocks', []))
        SimpleDocTemplate(str(path)).build(blocks)
    elif ext == '.csv':
        with path.open('w', encoding='utf-8-sig', newline='') as stream:
            csv.writer(stream).writerows(spec.get('rows', []))
    elif ext in TEXT:
        path.write_text(spec.get('text', ''), encoding='utf-8')
    else:
        raise ValueError('bb.error.unsupportedArtifact')
    return validate(path)


def pages(path, folder, start=0, limit=4):
    import fitz
    folder = Path(folder)
    folder.mkdir(parents=True, exist_ok=True)
    document = fitz.open(path)
    images = []
    for index in range(start, min(len(document), start + limit)):
        output = folder / f'page-{index + 1}.png'
        document[index].get_pixmap(matrix=fitz.Matrix(1.5, 1.5)).save(output)
        images.append(str(output))
    return {'pages': len(document), 'images': images, 'next': start + limit if start + limit < len(document) else None}


def main():
    args = json.loads(sys.stdin.read() or '{}')
    action = sys.argv[1]
    if action == 'validate':
        return validate(args['path'])
    if action == 'inspect':
        return inspect(args['path'])
    if action == 'render':
        return render(args['path'], args['folder'])
    if action == 'grid':
        return grid(args['path'], args.get('sheet'), max(0, int(args.get('offset', 0))), min(500, max(1, int(args.get('limit', 200)))))
    if action == 'create':
        return create(args['spec'], args['folder'])
    if action == 'pages':
        return pages(args['path'], args['folder'], max(0, int(args.get('start', 0))), min(8, max(1, int(args.get('limit', 4)))))
    raise ValueError('bb.error.unknownMethod')


if __name__ == '__main__':
    print(json.dumps(main(), ensure_ascii=False, default=str))
